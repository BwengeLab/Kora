import io
import importlib.util
import unittest
from pathlib import Path

from PIL import Image
from openpyxl import Workbook  # type: ignore[import-untyped]

from agents.document_ai.enterprise_schemas import BoundingBox, ContentLine, ContentPage, ContentToken
from agents.document_ai.raw_content import (
    _detect_page_tables,
    _group_positioned_lines,
    _positioned_line_text,
    _prepare_ocr_image,
    _split_line_into_cells,
    extract_raw_content,
)


class RawContentTests(unittest.TestCase):
    def test_groups_same_row_by_coordinates_without_interpreting_text(self) -> None:
        raw = [
            self.line("right", 500, 100),
            self.line("left", 100, 101),
            self.line("next", 100, 150),
        ]
        lines = _group_positioned_lines(raw, 1, 1000, 1000)
        self.assertEqual([line.text for line in lines], ["left\tright", "next"])
        self.assertEqual([token.text for token in lines[0].tokens], ["left", "right"])
        self.assertTrue(all(token.bounding_box is not None for token in lines[0].tokens))

    def test_large_word_gap_is_preserved_as_column_separator(self) -> None:
        raw = {
            "text": "$125 $1375",
            "words": [
                {"text": "$125", "x": 100, "y": 10, "width": 40, "height": 20},
                {"text": "$1375", "x": 156, "y": 10, "width": 50, "height": 20},
            ],
        }
        self.assertEqual(_positioned_line_text(raw), "$125\t$1375")

    def test_preprocessing_scales_small_document_without_changing_aspect_ratio(self) -> None:
        prepared = _prepare_ocr_image(Image.new("RGB", (1000, 500), "white"))
        self.assertEqual(prepared.size, (2000, 1000))

    def test_csv_preserves_exact_cells_and_quality(self) -> None:
        content = extract_raw_content(b"a,b\n1,2\n", "data.csv", "csv")
        self.assertEqual(content.pages[0].tables[0].row_count, 2)
        self.assertEqual(content.pages[0].tables[0].column_count, 2)
        self.assertEqual(
            [cell.text for cell in content.pages[0].tables[0].cells],
            ["a", "b", "1", "2"],
        )
        assert content.quality is not None
        self.assertEqual(content.quality.mean_confidence, 1.0)

    def test_excel_preserves_formulas_without_evaluating_business_meaning(self) -> None:
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["left", "right"])
        sheet.append([2, "=A2*3"])
        stream = io.BytesIO()
        workbook.save(stream)
        content = extract_raw_content(stream.getvalue(), "book.xlsx", "xlsx")
        table = content.pages[0].tables[0]
        self.assertEqual(table.row_count, 2)
        self.assertEqual([cell.text for cell in table.cells], ["left", "right", "2", "=A2*3"])

    @unittest.skipUnless(importlib.util.find_spec("pymupdf"), "PyMuPDF unavailable")
    def test_native_pdf_preserves_page_coordinates(self) -> None:
        root = Path(__file__).resolve().parents[2]
        content = extract_raw_content(
            (root / "testdata" / "synthetic" / "invoice_clean.pdf").read_bytes(),
            "document.pdf",
            "pdf",
        )
        self.assertEqual(len(content.pages), 1)
        self.assertTrue(content.full_text.strip())
        self.assertTrue(
            all(
                token.bounding_box is not None
                for line in content.pages[0].lines
                for token in line.tokens
            )
        )

    def test_repeated_geometric_rows_are_preserved_as_neutral_table(self) -> None:
        lines = []
        for row in range(3):
            tokens = tuple(
                ContentToken(
                    f"r{row}c{column}",
                    0.9,
                    BoundingBox(
                        1,
                        0.05 + column * 0.3,
                        0.1 + row * 0.1,
                        0.15 + column * 0.3,
                        0.14 + row * 0.1,
                    ),
                )
                for column in range(3)
            )
            lines.append(ContentLine(" ".join(token.text for token in tokens), row, tokens))
        tables = _detect_page_tables(ContentPage(1, 1000, 1000, tuple(lines)))
        self.assertEqual(len(tables), 1)
        self.assertEqual(tables[0].row_count, 3)
        self.assertEqual(tables[0].column_count, 3)

    def test_attached_table_border_starts_a_new_cell(self) -> None:
        tokens = (
            ContentToken("$88", 0.9, BoundingBox(1, 0.1, 0.1, 0.2, 0.2)),
            ContentToken("|$1,188", 0.9, BoundingBox(1, 0.25, 0.1, 0.4, 0.2)),
            ContentToken("tail", 0.9, BoundingBox(1, 0.7, 0.1, 0.8, 0.2)),
        )
        groups = _split_line_into_cells(tokens)
        self.assertEqual([[token.text for token in group] for group in groups], [["$88"], ["$1,188"], ["tail"]])

    @staticmethod
    def line(text: str, x: int, y: int) -> dict[str, object]:
        return {
            "text": text,
            "x": x,
            "y": y,
            "words": [{"text": text, "x": x, "y": y, "width": 40, "height": 20}],
        }


if __name__ == "__main__":
    unittest.main()
